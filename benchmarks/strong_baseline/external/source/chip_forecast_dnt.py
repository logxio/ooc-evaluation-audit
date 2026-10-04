#!/usr/bin/env python3
"""Third, independent benchmark: the frozen AnchorBoost on human neural cells (US EPA DNT screen, Harrill et al. 2018).

The US EPA tested 75 chemicals in a battery of high-content imaging and plate-reader assays for key events of brain
development, at up to 11 concentrations on three replicate plates (US EPA ScienceHub, DOI 10.23719/1407642, public
domain). This benchmark takes every endpoint measured in human cells: hNP1 neural progenitors (caspase apoptosis,
CellTiter viability and three proliferation readouts) and hN2 neurons (neuron count, neurite length, neurite count,
branch points), 9 endpoints for 71 chemicals. Rejected wells (quality 0) are removed; each plate is then centred on its
vehicle wells and each endpoint is divided by the robust SD of all its vehicle wells and clipped to +-10, as for the
acute MEA benchmark. A well is one plate position for one sample (one plate ID holds two chemicals in the source file),
and endpoints read from the same well share a row. Concentrations are grouped per chemical on a 0.1 log10 grid;
chemical folds come from a hash of the sample name. Model, features, hyperparameters, seeded designs, curve metric and
paired chemical bootstrap are those of chip_forecast.py, unchanged. No choice was made on this dataset, so all five
folds are held out (protocol: chip_forecast_dnt_protocol.json).
Run: python chip_forecast_dnt.py --folds 0 1 2 3 4 --out chip_forecast_dnt.json   (one fold per run also works;
     --merge dnt_fold*.csv summarizes per-fold CSVs)
"""
import argparse
import csv
import hashlib
import json
import urllib.request
from pathlib import Path

import numpy as np
import openpyxl

import chip_forecast as cf

ROOT = Path(__file__).resolve().parent
URL = 'https://pasteur.epa.gov/uploads/10.23719/1407642/Harrill%20et%20al%20DNT_Assay%20Dataset.xlsx'
SHA256 = 'ab6de83c38001c7d1cb977c3758bc77aa110d18658a3b9758369bf3a21d6bc0e'
HUMAN = ('hNP1', 'hN2')
K = 3
BASELINES = ('loglinear_interp', 'hill_per_endpoint', 'analog_knn')


def fetch(cache):
    path = cache / 'Harrill_DNT_Assay_Dataset.xlsx'
    if not path.exists():
        cache.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(URL, timeout=300) as r:
            path.write_bytes(r.read())
    got = hashlib.sha256(path.read_bytes()).hexdigest()
    if got != SHA256:
        raise ValueError(f'{path.name}: SHA-256 {got} differs from pinned {SHA256}')
    return path


def load_tasks(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = {r[0]: r[1] for r in list(wb['assay_component'].iter_rows(values_only=True))[1:]}
    aids = sorted(a for a in names if names[a].split('_')[0] in HUMAN)
    data = [r for r in list(wb['Data'].iter_rows(values_only=True))[1:] if r[1] in aids and r[7] == 1]
    wells, values = {}, []
    for j, a in enumerate(aids):
        rows = [r for r in data if r[1] == a]
        raw = np.array([np.nan if r[9] is None else float(r[9]) for r in rows])
        plate = np.array([r[3] for r in rows])
        vehicle = np.array([r[6] == 'n' for r in rows])
        v = raw.copy()
        for p in np.unique(plate):
            on = plate == p
            v[on] -= np.nanmedian(raw[on & vehicle])
        mad = 1.4826 * np.nanmedian(np.abs(v[vehicle]))
        v = np.clip(v / (mad if mad > 0 else np.nanstd(v[vehicle])), -10, 10)
        for r, x in zip(rows, v):
            if r[6] == 't' and np.isfinite(x):
                key = (r[3], r[4], r[5], r[2])
                wells.setdefault(key, (r[2], float(r[8])))
                values.append((key, j, x))
    index = {key: i for i, key in enumerate(wells)}
    y = np.zeros((len(index), 1, len(aids)), np.float32)
    m = np.zeros(y.shape, bool)
    for key, j, x in values:
        y[index[key], 0, j], m[index[key], 0, j] = x, True
    sample = np.array([wells[key][0] for key in index])
    logc = np.round(np.log10([wells[key][1] for key in index]), 1).astype(np.float32)
    cf.ND, cf.NF, cf.D = 1, len(aids), len(aids)  # one read time, 9 human-cell endpoints
    tasks = []
    for chem in sorted(set(sample)):
        i = np.where(sample == chem)[0]
        fold = int(hashlib.sha256(f'dnt-hci|{chem}'.encode()).hexdigest()[:8], 16) % 5
        tasks.append(cf.Task(chem, fold, 'unknown', logc[i], y[i], m[i]))
    return tasks, [names[a] for a in aids]


def summarize(rows, folds):
    pick = lambda name, fs: {r['chemical']: r['curve_mae'] for r in rows if r['method'] == name and r['fold'] in fs}
    ours = pick('anchorboost', folds)
    out = dict(n_chemicals=len(ours), anchorboost=round(float(np.mean(list(ours.values()))), 4), versus={}, fold_means={},
               fold_versus_loglinear={})
    for name in BASELINES:
        theirs = pick(name, folds)
        out[name] = round(float(np.mean(list(theirs.values()))), 4)
        out['versus'][name] = cf.paired(ours, theirs)
    for f in folds:
        out['fold_means'][f'fold{f}'] = {name: round(float(np.mean(list(pick(name, [f]).values()))), 4)
                                         for name in ('anchorboost',) + BASELINES}
        out['fold_versus_loglinear'][f'fold{f}'] = cf.paired(pick('anchorboost', [f]), pick('loglinear_interp', [f]))
    return out


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--cache', type=Path, default=ROOT / '.cache' / 'epa_dnt_hci')
    p.add_argument('--folds', type=int, nargs='+', default=[0, 1, 2, 3, 4])
    p.add_argument('--out', type=Path, default=ROOT / 'chip_forecast_dnt.json')
    p.add_argument('--merge', type=Path, nargs='+', help='summarize per-fold CSVs written by earlier runs')
    a = p.parse_args()
    if a.merge:
        rows = []
        for path in a.merge:
            with open(path, newline='') as f:
                rows += [dict(r, fold=int(r['fold']), curve_mae=float(r['curve_mae'])) for r in csv.DictReader(f)]
        folds, times = sorted({r['fold'] for r in rows}), None
    else:
        tasks, endpoints = load_tasks(fetch(a.cache))
        rows, times = cf.run(tasks, a.folds, [K])
        folds = a.folds
        with open(a.out.with_suffix('.csv'), 'w', newline='') as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
    res = dict(source=dict(url=URL, sha256=SHA256, cells=HUMAN), k=K, folds=folds, timing=times,
               summary=summarize(rows, folds), code_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
               model_sha256=hashlib.sha256(Path(cf.__file__).read_bytes()).hexdigest())
    a.out.write_text(json.dumps(res, indent=1) + '\n')
    print(json.dumps(res['summary'], indent=1))


if __name__ == '__main__':
    main()
