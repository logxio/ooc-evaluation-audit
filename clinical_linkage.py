#!/usr/bin/env python3
"""Audit the clinical-linkage denominator in a published on-chip lung organoid study.

This reanalyses the source workbook's eligibility markings. It does not infer
patient response from drug curves or treat unlinked lines as prediction errors.
"""

import argparse
import hashlib
import io
import json
import urllib.request
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


SOURCE_URL = ('https://static-content.springer.com/esm/'
              'art%3A10.1038%2Fs41467-021-22676-1/MediaObjects/'
              '41467_2021_22676_MOESM5_ESM.xlsx')
SOURCE_SHA256 = 'a9ba00bc14e3a4b0ecac6161243aff41ff192e1295622beeccd6003d105b91b7'


def source_bytes(path):
    if path:
        data = Path(path).read_bytes()
    else:
        with urllib.request.urlopen(SOURCE_URL, timeout=90) as response:
            data = response.read()
    digest = hashlib.sha256(data).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError(f'source SHA256 mismatch: {digest}')
    return data


def text(value):
    return '' if value is None else str(value).strip()


def is_author_linked(cell):
    fill = cell.fill
    return (fill.patternType == 'solid' and fill.fgColor.type == 'theme'
            and fill.fgColor.theme == 0
            and abs(fill.fgColor.tint - (-0.1499984740745262)) < 1e-9)


def unlinked_reason(record):
    reason = record['reason'].lower()
    if 'can not be evaluated' in reason or 'no subsequent treatment' in reason:
        return 'clinical_response_unavailable_or_no_treatment'
    if 'not tested' in reason or 'radiotherapy' in reason:
        return 'regimen_not_tested_or_confounded'
    return 'reason_unresolved'


def audit(data):
    sheet = load_workbook(io.BytesIO(data), data_only=True).active
    records = []
    for row in sheet.iter_rows(min_row=3):
        sample = text(row[0].value)
        if not sample or sample.startswith('*:'):
            continue
        if not (sample.startswith('LC') or sample.startswith('CRC')):
            continue
        record = {
            'sample_id': sample, 'source_row': row[0].row,
            'collection_time': text(row[1].value),
            'prior_treatment': text(row[3].value),
            'prior_response': text(row[4].value),
            'post_collection_treatment': text(row[5].value),
            'clinical_evaluability': text(row[7].value),
            'reported_response': text(row[8].value),
            'tested_agents': text(row[12].value),
            'reason': text(row[13].value),
            'author_linked': is_author_linked(row[0]),
        }
        if record['author_linked']:
            if record['reason']:
                raise ValueError(f'linked row has exclusion reason: {sample}')
            post = record['post_collection_treatment']
            record['has_post_collection_drug'] = post.lower() not in ('', 'none', 'radiotherapy')
            tested = {x.strip().lower() for x in record['tested_agents'].split(',')}
            record['exact_post_regimen_text_match'] = post.lower() in tested
            record['unlinked_reason'] = None
        else:
            record['has_post_collection_drug'] = False
            record['exact_post_regimen_text_match'] = False
            record['unlinked_reason'] = unlinked_reason(record)
        records.append(record)
    ids = [r['sample_id'] for r in records]
    if len(ids) != len(set(ids)):
        raise ValueError('duplicate Sample ID in source table')
    linked = [r for r in records if r['author_linked']]
    unlinked = [r for r in records if not r['author_linked']]
    reasons = Counter(r['unlinked_reason'] for r in unlinked)
    n = len(records)
    result = {
        'schema': 'clinical.linkage.coverage.v1',
        'source_article': 'https://doi.org/10.1038/s41467-021-22676-1',
        'source_workbook_url': SOURCE_URL,
        'source_workbook_sha256': SOURCE_SHA256,
        'source_license': 'CC BY 4.0',
        'tested_lines_in_source': n,
        'author_linked_lines': len(linked),
        'author_linked_coverage': round(len(linked) / n, 6),
        'unlinked_lines': len(unlinked),
        'unlinked_reason_counts': dict(sorted(reasons.items())),
        'linked_with_post_collection_drug': sum(r['has_post_collection_drug'] for r in linked),
        'linked_without_post_collection_drug': sum(not r['has_post_collection_drug'] for r in linked),
        'linked_with_exact_post_regimen_text_match': sum(r['exact_post_regimen_text_match'] for r in linked),
        'author_reported_agreements_among_linked': len(linked),
        'all_tested_agreement_bounds_if_unlinked_unknown': [len(linked), n],
        'limits': [
            'The original authors, not this audit, report ten of ten clinical agreements.',
            'Gray source rows include prior-treatment comparisons; a gray row is not a prospective prediction.',
            'Exact post-regimen text equality is descriptive and stricter than pharmacologic equivalence.',
            'No binary response threshold or independent clinical accuracy is recomputed here.',
            'Eleven unlinked lines are unknown, not incorrect or correct predictions.',
        ],
    }
    return result, records


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', help='Optional local copy of the original workbook')
    parser.add_argument('--out', help='Optional private patient-level audit JSON')
    args = parser.parse_args()
    result, records = audit(source_bytes(args.source))
    if args.out:
        Path(args.out).write_text(json.dumps({'summary': result, 'rows': records}, indent=2) + '\n')
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
