#!/usr/bin/env python3
"""Measure three concentrations, get the rest of the curve: forecasts, 90% intervals and the next step per compound.

Input is one long CSV with a row per well and endpoint:
    compound, concentration, unit, endpoint, value, plate, date
Vehicle-control wells, if any, have concentration 0. A row with an empty value names a concentration still to measure.

Forecast (any series with concentrations still to measure). A series with three measured concentrations, the highest
among them, gets a forecast and a 90% interval at every concentration still to measure and one next step: report the
call (active or inactive over the full series), add one named concentration and run this again with its readings, or
complete the series. With four measured concentrations the step is report or complete. Completed series of the same
endpoint, in the same table or in --library tables, train and calibrate the model; a compound being forecast never
trains its own endpoint. The output folder holds forecast.csv, summary.md, one model file per endpoint and lock.json
with the SHA-256 of input, output, models, protocol and code and the UTC time, to fix the forecast before the series
is finished. --reconcile then checks the finished series against the locked forecast, concentration by concentration.

Replay (every series complete). Each series is replayed from five three-concentration designs on five compound folds,
with training, calibration and test compounds kept apart, and compared with its full series: calls made, wrong calls,
wells used, forecast error and interval coverage.

The methods are the published ones of this repository, unchanged: AnchorBoost forecasts (chip_forecast.py), learned-
width 90% intervals (chip_forecast_intervals.py), the training-only cutoff, rank margin and group calibration with exact
arithmetic (paper_nested_decision.py), and the joint chain of the paper, which reports from three concentrations, adds
the concentration where AnchorBoost, interpolation and analog chemicals disagree most, or completes the series under one
margin calibrated over every candidate fourth concentration. Responses are put on the baseline scale of EPA's ToxCast
pipeline (tcpl): each plate is centred on its vehicle wells when it has them, then the median and robust SD (1.4826 x
MAD) of the responses at the two lowest concentrations of every completed compound set the baseline, computed from
training and calibration compounds only, and responses are clipped to +-10 baseline SDs. A series is active when its
largest absolute concentration mean reaches 3 baseline SDs (--threshold).

Run: python three_point.py my_table.csv                       (forecast if any series is unfinished, else replay)
     python three_point.py my_table.csv --library more.csv    (more completed series of the same endpoints)
     python three_point.py my_finished_table.csv --reconcile three_point_my_table
     python three_point.py --example epa_dnt                  (the US EPA human neural screen in this format)
"""
import argparse
import csv
import hashlib
import importlib.util
import json
import os
import pickle
import platform
import re
import sys
import time
import types
from collections import defaultdict
from concurrent.futures import ProcessPoolExecutor
from datetime import datetime, timezone
from fractions import Fraction
from pathlib import Path

import numpy as np
import sklearn
from sklearn.ensemble import HistGradientBoostingRegressor
from threadpoolctl import threadpool_limits

import chip_forecast as cf
import chip_forecast_intervals as ci
import chip_forecast_selfcheck as sc

ROOT = Path(__file__).resolve().parent
K = 3
CLIP = 10.0
TOL = 0.05  # log10 units: concentrations within ~12% of each other are one level
PARTS = 4  # calibration parts for the cross-fitted interval width models
COLUMNS = {'compound': ('compound', 'chemical', 'drug', 'treatment', 'sample'),
           'concentration': ('concentration', 'conc', 'dose'),
           'unit': ('unit', 'units', 'concentration_unit'),
           'endpoint': ('endpoint', 'assay', 'readout'),
           'value': ('value', 'reading', 'response', 'raw_value'),
           'plate': ('plate', 'replicate', 'plate_id', 'replicate_or_plate'),
           'date': ('date', 'culture_date', 'run_date')}
REQUIRED = ('compound', 'concentration', 'endpoint', 'value')
MOLAR = {'pm': 1e-6, 'nm': 1e-3, 'um': 1.0, 'mm': 1e3, 'm': 1e6}
EMPTY = {'', 'na', 'nan', 'none', 'null', '-'}
CODE = ['three_point.py', 'chip_forecast.py', 'chip_forecast_intervals.py', 'chip_forecast_selfcheck.py',
        'paper_nested_decision.py']


def rule_module():
    """paper_nested_decision.py imports pandas, psutil and rdata at module level for its EPA preparation stages. This
    entry calls only its release-rule functions, so each of the three that is not installed stands in during that
    import as a placeholder that raises if touched; sys.modules is restored afterwards."""
    class Missing(types.ModuleType):
        def __getattr__(self, name):
            raise ImportError(f'{self.__name__} is needed only by the EPA preparation stages of paper_nested_decision.py')
    added = [n for n in ('pandas', 'psutil', 'rdata') if importlib.util.find_spec(n) is None]
    for n in added:
        sys.modules[n] = Missing(n)
    try:
        import paper_nested_decision
    finally:
        for n in added:
            sys.modules.pop(n, None)
    return paper_nested_decision


nd = rule_module()


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def deal(names, key, n):
    """Balanced deterministic partition: names sorted by SHA-256 of (key, name), dealt round-robin into n parts."""
    order = sorted(names, key=lambda c: nd.digest([key, c]))
    return {c: i % n for i, c in enumerate(order)}


def slug(text):
    return re.sub(r'[^A-Za-z0-9._-]+', '_', text).strip('_') or 'endpoint'


def runtime(start):
    """Wall seconds and peak resident memory of this process and of its largest worker process."""
    import resource
    unit = 1 if sys.platform == 'darwin' else 1024
    main = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * unit
    worker = resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss * unit
    return dict(seconds=round(time.time() - start, 1), peak_rss_mb=round(main / 2 ** 20),
                largest_worker_peak_rss_mb=round(worker / 2 ** 20))


# ---------------------------------------------------------------- reading a laboratory table

def read_table(path):
    """Rows of a long laboratory CSV with canonical keys; value None marks a concentration still to measure."""
    with open(path, newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        head = {h.strip().lower(): h for h in (reader.fieldnames or [])}
        pick = {}
        for key, names in COLUMNS.items():
            pick[key] = next((head[n] for n in names if n in head), None)
            if pick[key] is None and key in REQUIRED:
                raise ValueError(f'{path}: no {key} column (accepted names: {", ".join(names)})')
        rows = []
        for line, r in enumerate(reader, 2):
            get = lambda k: (r.get(pick[k]) or '').strip() if pick[k] else ''
            compound, endpoint = get('compound'), get('endpoint')
            if not compound or not endpoint:
                raise ValueError(f'{path} line {line}: compound and endpoint are required')
            try:
                conc = float(get('concentration'))
            except ValueError:
                raise ValueError(f'{path} line {line}: concentration {get("concentration")!r} is not a number') from None
            if not conc >= 0:
                raise ValueError(f'{path} line {line}: concentration must be 0 (vehicle) or positive')
            raw = get('value')
            try:
                value = None if raw.lower() in EMPTY else float(raw)
            except ValueError:
                raise ValueError(f'{path} line {line}: value {raw!r} is not a number') from None
            if value is not None and not np.isfinite(value):
                value = None
            unit = get('unit')
            u = unit.lower().replace('µ', 'u').replace('μ', 'u').replace(' ', '')
            if conc > 0 and u in MOLAR:
                conc, unit = conc * MOLAR[u], 'uM'
            rows.append(dict(compound=compound, conc=conc, unit=unit or 'as given', endpoint=endpoint, value=value,
                             plate=(get('plate'), get('date')), line=line))
    if not rows:
        raise ValueError(f'{path}: no rows')
    return rows


def plate_centres(rows):
    """Per endpoint: the median of each plate's vehicle wells (plate, date), and the median of those for plates without
    vehicle wells; None for an endpoint without vehicle wells, whose values are then used as read."""
    out = {}
    for ep in sorted({r['endpoint'] for r in rows}):
        by = defaultdict(list)
        for r in rows:
            if r['endpoint'] == ep and r['conc'] == 0 and r['value'] is not None:
                by[r['plate']].append(r['value'])
        centre = {p: float(np.median(v)) for p, v in by.items()}
        out[ep] = dict(centre=centre, fallback=float(np.median(list(centre.values()))),
                       vehicle_wells=sum(map(len, by.values()))) if centre else None
    return out


def level_values(logs):
    """One representative log10 concentration per value: values within TOL of a neighbour form one level (median)."""
    logs = np.asarray(logs, float)
    order, rep, start = np.argsort(logs), np.empty(len(logs)), 0
    for i in range(1, len(order) + 1):
        if i == len(order) or logs[order[i]] - logs[order[i - 1]] > TOL:
            idx = order[start:i]
            rep[idx] = np.median(logs[idx])
            start = i
    return np.round(rep, 4).astype(np.float32)


class Curve:
    """One compound and endpoint of one table: measured wells centred on their plate's vehicle wells, and the
    concentrations still to measure."""

    def __init__(self, compound, endpoint, rows, centres, source):
        self.compound, self.endpoint, self.source = compound, endpoint, source
        units = sorted({r['unit'] for r in rows})
        if len(units) > 1:
            raise ValueError(f'{source}: {compound} / {endpoint} mixes concentration units {units}')
        self.unit = units[0]
        logs = level_values(np.log10([r['conc'] for r in rows]))
        self.display = {float(v): float(np.median([r['conc'] for r, l in zip(rows, logs) if l == v])) for v in set(logs.tolist())}
        meas = [i for i, r in enumerate(rows) if r['value'] is not None]
        self.logc = logs[meas]
        self.values = np.array([rows[i]['value'] for i in meas], float)
        base = np.array([centres['centre'].get(rows[i]['plate'], centres['fallback']) if centres else 0.0 for i in meas])
        self.centred = self.values - base
        self.centre = float(np.median(base)) if len(base) else 0.0
        self.levels = np.unique(self.logc)
        self.planned = np.array(sorted(set(logs.tolist()) - set(self.logc.tolist())), np.float32)
        self.open = len(self.planned) > 0

    def task(self, scale):
        """chip_forecast.Task in baseline SDs, clipped to +-10."""
        z = np.clip((self.centred - scale['b0']) / scale['b1'], -CLIP, CLIP).astype(np.float32)
        t = cf.Task(self.compound, 0, 'unused', self.logc, z.reshape(-1, 1, 1), np.ones((len(z), 1, 1), bool))
        t.group = self.compound
        return t

    def raw(self, z, scale):
        return self.centre + scale['b0'] + np.asarray(z, float) * scale['b1']

    def measured_mean(self, level):
        return float(self.values[self.logc == level].mean())


def baseline(curves):
    """tcpl baseline: median and 1.4826 x MAD of the centred responses at the two lowest concentrations of every series."""
    v = np.concatenate([c.centred[np.isin(c.logc, c.levels[:2])] for c in curves])
    b0 = float(np.median(v))
    b1 = 1.4826 * float(np.median(np.abs(v - b0)))
    if not b1 > 0:
        b1 = float(np.std(v, ddof=1)) if len(v) > 1 else 0.0
    if not b1 > 0:
        raise ValueError(f'{curves[0].endpoint}: responses at the lowest concentrations do not vary, so no baseline scale exists')
    return dict(b0=b0, b1=b1, wells=int(len(v)), series=len(curves))


def load_curves(path, endpoints=None):
    rows = read_table(path)
    if endpoints:
        rows = [r for r in rows if r['endpoint'] in endpoints]
        if not rows:
            raise ValueError(f'{path}: none of the endpoints {sorted(endpoints)} is in this table')
    centres = plate_centres(rows)
    by = defaultdict(list)
    for r in rows:
        if r['conc'] > 0:
            by[(r['compound'], r['endpoint'])].append(r)
    return [Curve(c, e, rr, centres[e], str(path)) for (c, e), rr in sorted(by.items())]


def truth(task, threshold):
    """1 = inactive over the full series (largest absolute concentration mean below the threshold), 0 = active."""
    return int(nd.effect(task.mu, task.ok) < threshold)


def observed(task, levels):
    ix = np.isin(task.logc, levels)
    return cf.Task(task.chem, task.fold, 'unused', task.logc[ix], task.y[ix], task.m[ix])


def score(model, obs, query):
    """Largest absolute response over the series: measured concentrations as measured, the others as forecast."""
    measured = nd.effect(obs.mu, obs.ok)
    if not len(query):
        return measured, np.zeros((0, cf.ND, cf.NF), np.float32)
    pred = model(obs, np.ones(len(obs.logc), bool), query)
    return max(measured, float(np.abs(pred.mean(axis=1)).max())), pred


def predict_many(model, pairs):
    """AnchorBoost forecasts for many (measured series, query levels) pairs in one tree-ensemble call; per row the same
    arithmetic as chip_forecast.AnchorBoost.__call__."""
    bases, feats = [], []
    for obs, query in pairs:
        base, f = cf.features(obs, np.ones(len(obs.logc), bool), query, model.analog, model.profiles, model.ablation)
        bases.append(base)
        feats.append(f.reshape(-1, f.shape[-1]))
    pred = model.model.predict(np.concatenate(feats))
    out, at = [], 0
    for (obs, query), base in zip(pairs, bases):
        n = len(query) * cf.D
        out.append((pred[at:at + n].reshape(len(query), cf.D) + base * (model.ablation != 'no_anchor'))
                   .reshape(len(query), cf.ND, cf.NF))
        at += n
    return out


def disagreement(model, obs, query, pred):
    """The paper's model_disagreement choice: spread of AnchorBoost, interpolation and analog forecasts; ties go lower."""
    ones = np.ones(len(obs.logc), bool)
    stack = np.stack([pred, cf.predict_interp(obs, ones, query), model.analog.predict(obs, ones, query)])
    return int(np.argmax(np.std(stack, axis=0).mean(axis=(1, 2))))


# ---------------------------------------------------------------- fitting one endpoint

def reveal(records, points, tasks, threshold):
    by = {t.chem: t for t in tasks}
    for r in records:
        r['truth'] = truth(by[r['chemical']], threshold)
    points['truth'] = np.array([by[str(c)].mu[l] for c, l in zip(points['chemical'], points['level'])])
    points['mask'] = np.array([by[str(c)].ok[l] for c, l in zip(points['chemical'], points['level'])])


def chain_rows(model, rule, tasks, records):
    """For every design (the paper's joint chain): the three-concentration call and margin, the disagreement choice,
    and for every candidate fourth concentration the call and margin from the four measured ones."""
    by = {(r['chemical'], r['design']): r for r in records}
    rows = []
    for t in tasks:
        ctxs = cf.designs(t, K)
        keys = [tuple(c) for c in ctxs]
        for ctx in ctxs:
            keys += [tuple(np.sort(np.r_[ctx, level])) for level in t.levels[~np.isin(t.levels, ctx)]]
        keys = [k for k in dict.fromkeys(keys) if len(k) < len(t.levels)]
        obs = {k: observed(t, np.array(k, np.float32)) for k in keys}
        query = {k: t.levels[~np.isin(t.levels, np.array(k, np.float32))] for k in keys}
        preds = dict(zip(keys, predict_many(model, [(obs[k], query[k]) for k in keys])))

        def scored(k):
            o = obs[k] if k in obs else observed(t, np.array(k, np.float32))
            measured = nd.effect(o.mu, o.ok)
            return max(measured, float(np.abs(preds[k].mean(axis=1)).max())) if k in preds else measured

        for di, ctx in enumerate(ctxs):
            k3, cand = tuple(ctx), t.levels[~np.isin(t.levels, ctx)]
            s3 = scored(k3)
            if abs(s3 - by[(t.chem, di)]['anchorboost']) > 1e-6:
                raise AssertionError(f'{t.chem} design {di}: chain score differs from the release-rule score')
            call3, margin3 = nd.apply_rule(rule, s3)
            branches = []
            for level in cand:
                call4, margin4 = nd.apply_rule(rule, scored(tuple(np.sort(np.r_[ctx, level]))))
                branches.append(dict(level=float(level), wells=int(np.sum(t.logc == level)), call4=call4, margin4=margin4))
            rows.append(dict(compound=t.chem, group=t.group, design=di, context=ctx, call3=call3, margin3=margin3,
                             choice=disagreement(model, obs[k3], cand, preds[k3]), branches=branches,
                             wells_initial=int(np.isin(t.logc, ctx).sum()), wells_full=len(t.logc)))
    return rows


def chain_margin(rows, alpha):
    """The paper's joint calibration with exact arithmetic: a design's loss at margin m is 1 when the three-concentration call
    or any candidate four-concentration call clears m and is wrong; group means; the first grid margin with
    (sum + 1) / (groups + 1) <= alpha, else 1 (every design is completed)."""
    groups = defaultdict(list)
    for r in rows:
        groups[r['group']].append(r)
    level, n = Fraction(repr(alpha)), len(groups)
    for margin in nd.MARGIN_GRID:
        total = Fraction(0)
        for rr in groups.values():
            bad = sum(max([int(r['margin3'] > margin and r['call3'] != r['truth'])] +
                          [int(b['margin4'] > margin and b['call4'] != r['truth']) for b in r['branches']]) for r in rr)
            total += Fraction(bad, len(rr))
        if (total + 1) / (n + 1) <= level:
            return margin, True
    return nd.MARGIN_GRID[-1], False


def well_errors(tasks, points):
    """Interval covariates (chip_forecast_intervals.covariates) and absolute errors of every held-out well, per design."""
    pred = {(str(c), int(d), int(l)): float(np.asarray(p).ravel()[0])
            for c, d, l, p in zip(points['chemical'], points['design'], points['level'], points['anchorboost'])}
    out = {}
    for t in tasks:
        X, e = [], []
        for di, ctx in enumerate(cf.designs(t, K)):
            held = np.where(~np.isin(t.logc, ctx))[0]
            W = np.nonzero(t.m[held].reshape(len(held), cf.D))[0]
            lev = np.searchsorted(t.levels, t.logc[held[W]])
            X.append(ci.covariates(t, ctx)[0])
            e.append(np.abs(np.array([pred[(t.chem, di, int(l))] for l in lev]) - t.y[held[W]].reshape(-1)))
        out[t.chem] = (np.concatenate(X), np.concatenate(e).astype(np.float32))
    return out


def width_models(errors, key, rng):
    """chip_forecast_intervals.py widths: one quantile model per calibration part, trained on the other parts; the
    conformal quantile of errors scaled by the model that did not see them."""
    part = deal(sorted(errors), ['width-v1', key], PARTS)
    parts = {g: [errors[c] for c in errors if part[c] == g] for g in range(PARTS)}
    parts = {g: (np.concatenate([x for x, _ in v]), np.concatenate([e for _, e in v])) for g, v in parts.items() if v}
    if len(parts) < 2:
        raise ValueError('interval widths need at least two calibration compounds')
    models, scores = [], []
    for g in sorted(parts):
        X = np.concatenate([parts[h][0] for h in parts if h != g])
        e = np.concatenate([parts[h][1] for h in parts if h != g]).astype(np.float64)
        idx = rng.choice(len(X), min(ci.ROWS, len(X)), replace=False)
        m = HistGradientBoostingRegressor(**ci.WIDTH).fit(X[idx], e[idx])
        models.append(m)
        scores.append(parts[g][1] / np.maximum(m.predict(parts[g][0]), ci.FLOOR))
    return dict(models=models, q=float(sc.quantile(np.concatenate(scores))))


def half_width(widths, X):
    return widths['q'] * np.mean([np.maximum(m.predict(X), ci.FLOOR) for m in widths['models']], axis=0)


def fit(train, cal, key, threshold, alpha):
    """Everything a decision needs, from training and calibration compounds only (paper_nested_decision.py order):
    training-only cutoffs and rank margins from two inner fits, the final AnchorBoost on all training compounds, then on
    calibration compounds the chain margin, the two-action margins and the interval width models."""
    if len(train) < 2 or len(cal) < 2:
        raise ValueError(f'needs at least 2 training and 2 calibration compounds; has {len(train)} and {len(cal)}')
    half = deal([t.chem for t in train], ['inner-v1', key], 2)
    records, points = [], []
    for h in (0, 1):
        held = [t for t in train if half[t.chem] == h]
        rec, pts = nd.forecast(cf.AnchorBoost([t for t in train if half[t.chem] != h], K), held, key, f'inner{h}')
        reveal(rec, pts, held, threshold)
        records += rec
        points.append(pts)
    rules = nd.fit_rules(records, points, {t.chem: {'drug_group': t.group} for t in train})
    model = cf.AnchorBoost(train, K)
    cal_rec, cal_pts = nd.forecast(model, cal, key, 'calibration')
    chain = chain_rows(model, rules['methods']['anchorboost'], cal, cal_rec)
    cal_rec = nd.apply_rules(cal_rec, rules)
    reveal(cal_rec, cal_pts, cal, threshold)
    for r in chain:
        r['truth'] = truth(next(t for t in cal if t.chem == r['compound']), threshold)
    margin, feasible = chain_margin(chain, alpha)
    two = {m: nd.calibrate(cal_rec, m, alpha)['margin'] for m in ('anchorboost', 'measured_only')}
    widths = width_models(well_errors(cal, cal_pts), key, np.random.default_rng(0))
    return dict(model=model, rules=rules, chain_margin=margin, chain_feasible=feasible, two_action=two, widths=widths,
                training=sorted(t.chem for t in train), calibration=sorted(t.chem for t in cal), threshold=threshold,
                alpha=alpha, training_rows=model.rows)


# ---------------------------------------------------------------- forecast mode

def decide_open(f, t, q):
    """Forecast, 90% well intervals and the next step for one series with concentrations q still to measure."""
    k = len(t.levels)
    out = dict(action='', call='', add=None, margin=None, pred=None, half=None)
    if k < K:
        out['action'] = 'measure three concentrations first, the highest among them'
        return out
    s, pred = score(f['model'], t, q)
    call, margin = nd.apply_rule(f['rules']['methods']['anchorboost'], s)
    pseudo = cf.Task(t.chem, 0, 'unused', np.r_[t.logc, q].astype(np.float32),
                     np.r_[t.y, np.zeros((len(q), 1, 1), np.float32)], np.r_[t.m, np.ones((len(q), 1, 1), bool)])
    out.update(call='inactive' if call else 'active', margin=float(margin), pred=pred.reshape(-1),
               half=half_width(f['widths'], ci.covariates(pseudo, t.levels)[0]))
    if margin > f['chain_margin']:
        out['action'] = 'report the call'
    elif k == K and len(q) > 1:
        out['add'] = float(q[disagreement(f['model'], t, q, pred)])
        out['action'] = 'add one concentration'
    else:
        out['action'] = 'complete the series'
    return out


def forecast_endpoint(ep, curves, library, threshold, alpha, out):
    """Fit one endpoint on its completed series, save the fitted model, and forecast its unfinished series."""
    forecast_names = {c.compound for c in curves}
    lib = [c for c in library if c.endpoint == ep and c.compound not in forecast_names]
    dropped = sorted({c.compound for c in library if c.endpoint == ep and c.compound in forecast_names})
    notes = [f'{ep}: completed series of {", ".join(dropped)} left out of training, since those compounds are being '
             'forecast.'] if dropped else []
    units = {c.unit for c in lib + curves}
    if len(units) > 1:
        raise ValueError(f'endpoint {ep!r}: concentrations in different units {sorted(units)} cannot share a model')
    usable = [c for c in lib if len(c.levels) > K]
    if len(usable) < 4:
        raise ValueError(f'endpoint {ep!r}: {len(usable)} completed series with at least four concentrations; '
                         'add completed compounds of this endpoint to the table or with --library')
    scale = baseline(usable)
    role = deal([c.compound for c in usable], ['library-v1', ep], 2)
    f = fit([c.task(scale) for c in usable if role[c.compound] == 0],
            [c.task(scale) for c in usable if role[c.compound] == 1], f'A|{ep}', threshold, alpha)
    f['scale'] = scale
    path = out / f'model_{slug(ep)}.pkl'
    with open(path, 'wb') as fh:
        pickle.dump(dict(endpoint=ep, fit=f), fh, protocol=5)
    model = dict(file=path.name, sha256=sha(path))
    settings = dict(training_compounds=len(f['training']), calibration_compounds=len(f['calibration']),
                    training_rows=f['training_rows'], release_margin=f['chain_margin'],
                    calibration_feasible=f['chain_feasible'], interval_quantile=round(f['widths']['q'], 4),
                    cutoff=round(f['rules']['methods']['anchorboost']['cutoff'], 4),
                    baseline=dict(median=round(scale['b0'], 6), sd=round(scale['b1'], 6), wells=scale['wells']))
    table = []
    for c in curves:
        t = c.task(scale)
        d = decide_open(f, t, c.planned)
        for lv in np.unique(np.r_[c.levels, c.planned]):
            row = dict(compound=c.compound, endpoint=ep, concentration=c.display[float(lv)], unit=c.unit)
            i = int(np.searchsorted(t.levels, lv))
            if i < len(t.levels) and t.levels[i] == lv:
                row.update(status='measured', wells=int(np.sum(t.logc == lv)), value=round(c.measured_mean(lv), 6),
                           low90='', high90='', response=round(float(t.mu[i].ravel()[0]), 4), response_low90='',
                           response_high90='')
            elif d['pred'] is not None:
                j = int(np.searchsorted(c.planned, lv))
                z, h = float(d['pred'][j]), float(d['half'][j])
                row.update(status='forecast', wells='', value=round(float(c.raw(z, scale)), 6),
                           low90=round(float(c.raw(z - h, scale)), 6), high90=round(float(c.raw(z + h, scale)), 6),
                           response=round(z, 4), response_low90=round(z - h, 4), response_high90=round(z + h, 4))
            else:
                row.update(status='to measure', wells='', value='', low90='', high90='', response='',
                           response_low90='', response_high90='')
            row.update(next_step=d['action'], call=d['call'],
                       add_concentration='' if d['add'] is None else c.display[d['add']],
                       call_margin='' if d['margin'] is None else round(d['margin'], 6),
                       release_margin=f['chain_margin'], measured_concentrations=len(c.levels))
            table.append(row)
    return dict(endpoint=ep, table=table, model=model, notes=notes, settings=settings)


def worker_init(threads):
    """Each worker process: one output per series, and its share of the CPU threads."""
    global LIMITS
    cf.ND, cf.NF, cf.D = 1, 1, 1
    LIMITS = threadpool_limits(limits=threads)


def run_all(jobs, fn, tasks):
    """Run fn over argument tuples, in a pool of fresh worker processes (one task each) when jobs > 1. Tree fitting is
    mostly single-threaded Python, so workers share the cores rather than each worker taking several threads."""
    if jobs <= 1 or len(tasks) <= 1:
        return [fn(*t) for t in tasks]
    jobs = min(jobs, len(tasks))
    threads = max(1, (os.cpu_count() or 2) // jobs)
    os.environ['OMP_NUM_THREADS'] = str(threads)
    with ProcessPoolExecutor(max_workers=jobs, max_tasks_per_child=1, initializer=worker_init,
                             initargs=(threads,)) as ex:
        return [x.result() for x in [ex.submit(fn, *t) for t in tasks]]


def summary_forecast(table, settings, notes, args, lock_fields):
    curves = defaultdict(list)
    for r in table:
        curves[(r['compound'], r['endpoint'])].append(r)
    groups = defaultdict(list)
    for (comp, ep), rr in sorted(curves.items()):
        meas = ', '.join(f'{r["concentration"]:g}' for r in rr if r['status'] == 'measured') + f' {rr[0]["unit"]}'
        groups[rr[0]['next_step']].append((comp, ep, meas, rr[0]))
    n = {k: len(v) for k, v in groups.items()}
    need = lambda k: f'{k} need{"s" if k == 1 else ""}'
    parts = [f'{n.get("report the call", 0)} can be called now', f'{need(n.get("add one concentration", 0))} one more concentration',
             f'{need(n.get("complete the series", 0))} the full series']
    if n.get('measure three concentrations first, the highest among them'):
        parts.append(f'{need(n["measure three concentrations first, the highest among them"])} three measured concentrations first')
    lines = [f'# Next steps for {Path(args.table).name}', '',
             f'{len(curves)} series: ' + ', '.join(parts) + '. '
             f'Fixed at {lock_fields["created_utc"]} (UTC); send lock.json, or the SHA-256 the command prints, to the '
             'laboratory before the remaining concentrations are read. `forecast.csv` holds the forecast and the 90% '
             'interval for one well at every concentration still to measure; `--reconcile` checks them against the '
             'finished series concentration by concentration.']
    sections = [('report the call', 'Call now', ('Compound', 'Endpoint', 'Call', 'Measured'),
                 lambda c, e, m, r: (c, e, r['call'], m)),
                ('add one concentration', 'Add one concentration, then run again with its readings',
                 ('Compound', 'Endpoint', 'Add', 'Measured'),
                 lambda c, e, m, r: (c, e, f'{r["add_concentration"]:g} {r["unit"]}', m)),
                ('complete the series', 'Complete the series', ('Compound', 'Endpoint', 'Measured'),
                 lambda c, e, m, r: (c, e, m)),
                ('measure three concentrations first, the highest among them', 'Measure three concentrations first, the highest among them',
                 ('Compound', 'Endpoint', 'Measured'), lambda c, e, m, r: (c, e, m))]
    for key, title, head, cells in sections:
        if groups.get(key):
            lines += ['', f'## {title}', '', '| ' + ' | '.join(head) + ' |', '|' + '|'.join(':--' for _ in head) + '|']
            lines += ['| ' + ' | '.join(str(x) for x in cells(*g)) + ' |' for g in groups[key]]
    lines += ['', '**How the next step is chosen.** Values are in the table\'s own units. The model works in baseline '
              'SDs: the spread of responses at the two lowest concentrations of the completed compounds, after centring '
              'each plate on its vehicle wells when the table has them. A call is '
              f'active when the largest absolute concentration mean of the full series reaches {args.threshold:g} baseline '
              'SDs. The call is reported when its rank margin clears the release '
              f'margin, set on calibration compounds so that the expected share of wrong early calls per compound stays '
              f'at or below {args.alpha:.0%} whichever fourth concentration is added; otherwise the concentration where '
              'the forecasters disagree most is added, and after four measured concentrations the series is reported '
              'or completed.', '']
    for ep, s in sorted(settings.items()):
        lines.append(f'- {ep}: trained on {s["training_compounds"]} completed compounds ({s["training_rows"]:,} '
                     f'design rows), calibrated on {s["calibration_compounds"]}; baseline {s["baseline"]["median"]:.4g} '
                     f'± {s["baseline"]["sd"]:.4g} from {s["baseline"]["wells"]:,} wells; release margin {s["release_margin"]:g}'
                     + ('' if s['calibration_feasible'] else ' (too few calibration compounds for any early call at '
                        'this level, so every series is completed)') + '.')
    lines += [f'- {n}' for n in notes]
    lines += ['', '**Lock.** SHA-256 of what produced this forecast:', '']
    lines += [f'- {k}: `{v}`' for k, v in lock_fields['digest_lines']]
    return '\n'.join(lines) + '\n'


# ---------------------------------------------------------------- replay mode

def decide_chain(r, margin):
    """The paper's joint chain for one test design at the calibrated margin (model_disagreement policy)."""
    if r['margin3'] > margin:
        return dict(step='report at three', call=r['call3'], added=None, wells_added=0, complete=0)
    b = r['branches'][r['choice']]
    if b['margin4'] > margin:
        return dict(step='report after one more', call=b['call4'], added=b['level'], wells_added=b['wells'], complete=0)
    return dict(step='complete', call=None, added=b['level'], wells_added=b['wells'], complete=1)


def replay_fold(curves, ep, f, threshold, alpha):
    """One outer fold of one endpoint: fit on training and calibration compounds, decide every test design, then read
    the test compounds' full series to score the decisions, forecasts and intervals."""
    unit = curves[0].unit
    fold = deal([c.compound for c in curves], ['outer-v1', ep], 5)
    other = [g for g in range(5) if g != f]
    # The baseline scale comes from training and calibration compounds; test compounds only receive it.
    scale = baseline([c for c in curves if fold[c.compound] != f])
    tasks = {c.compound: c.task(scale) for c in curves}
    test = [t for t in tasks.values() if fold[t.chem] == f]
    fitted = fit([t for t in tasks.values() if fold[t.chem] in other[:2]],
                 [t for t in tasks.values() if fold[t.chem] in other[2:]], f'B|{ep}|{f}', threshold, alpha)
    rec, pts = nd.forecast(fitted['model'], test, f, 'test')
    chain = chain_rows(fitted['model'], fitted['rules']['methods']['anchorboost'], test, rec)
    rec = nd.apply_rules(rec, fitted['rules'])
    # Full series of the test compounds are read only from here on, after every forecast and decision is made.
    reveal(rec, pts, test, threshold)
    by = {(r['chemical'], r['design']): r for r in rec}
    rows = []
    for r in chain:
        t = tasks[r['compound']]
        y = truth(t, threshold)
        d = decide_chain(r, fitted['chain_margin'])
        n = by[(r['compound'], r['design'])]
        ab = n['anchorboost_margin'] > fitted['two_action']['anchorboost']
        mo = n['measured_only_margin'] > fitted['two_action']['measured_only']
        used = r['wells_initial'] + d['wells_added'] + d['complete'] * (r['wells_full'] - r['wells_initial'] - d['wells_added'])
        rows.append(dict(endpoint=ep, compound=r['compound'], fold=f, design=r['design'],
                         measured=';'.join(f'{10 ** float(v):.4g}' for v in r['context']), unit=unit,
                         includes_highest=int(np.isclose(r['context'].max(), t.levels.max())),
                         full_series_call='inactive' if y else 'active', next_step=d['step'],
                         call='' if d['call'] is None else 'inactive' if d['call'] else 'active',
                         added='' if d['added'] is None else f'{10 ** d["added"]:.4g}',
                         wrong=int(d['call'] is not None and d['call'] != y),
                         wells_initial=r['wells_initial'], wells_added=d['wells_added'], wells_used=used,
                         wells_full=r['wells_full'], margin3=round(float(r['margin3']), 6),
                         release_margin=fitted['chain_margin'],
                         two_action_report=int(ab), two_action_wrong=int(ab and n['anchorboost_call'] != y),
                         two_action_wells=r['wells_initial'] if ab else r['wells_full'],
                         measured_only_report=int(mo), measured_only_wrong=int(mo and n['measured_only_call'] != y),
                         measured_only_wells=r['wells_initial'] if mo else r['wells_full']))
    mae = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    for c, d_, p, lp, tr, m in zip(pts['chemical'], pts['design'], pts['anchorboost'], pts['loglinear'], pts['truth'], pts['mask']):
        if m.any():
            mae['anchorboost'][str(c)][int(d_)].append(float(np.abs(p - tr)[m].mean()))
            mae['loglinear'][str(c)][int(d_)].append(float(np.abs(lp - tr)[m].mean()))
    curve = {m: {c: float(np.mean([np.mean(v) for v in by_design.values()])) for c, by_design in mae[m].items()}
             for m in ('anchorboost', 'loglinear')}
    errors = well_errors(test, pts)
    hits, widths = [], []
    for t in test:
        h = half_width(fitted['widths'], errors[t.chem][0])
        hits.append(errors[t.chem][1] <= h)
        widths.append(h)
    info = dict(training=len(fitted['training']), calibration=len(fitted['calibration']), test=len(test),
                baseline=dict(median=round(scale['b0'], 6), sd=round(scale['b1'], 6)), release_margin=fitted['chain_margin'],
                calibration_feasible=fitted['chain_feasible'], two_action_margins=fitted['two_action'],
                interval_quantile=round(fitted['widths']['q'], 4))
    return dict(endpoint=ep, fold=f, rows=rows, curve=curve, hits=np.concatenate(hits), widths=np.concatenate(widths), info=info)


def replay_summary(parts):
    """Per-endpoint totals over the outer folds."""
    curve = {m: {c: v for p in parts for c, v in p['curve'][m].items()} for m in ('anchorboost', 'loglinear')}
    hits = np.concatenate([p['hits'] for p in parts])
    widths = np.concatenate([p['widths'] for p in parts])
    rows = [r for p in sorted(parts, key=lambda p: p['fold']) for r in p['rows']]
    top = [r for r in rows if r['includes_highest']]
    return rows, dict(folds={f'fold{p["fold"]}': p['info'] for p in sorted(parts, key=lambda p: p['fold'])},
                      curve_mae=dict(anchorboost=round(float(np.mean(list(curve['anchorboost'].values()))), 4),
                                     loglinear=round(float(np.mean(list(curve['loglinear'].values()))), 4),
                                     paired=cf.paired(curve['anchorboost'], curve['loglinear'])),
                      intervals=dict(wells=int(len(hits)), coverage=round(float(hits.mean()), 4),
                                     mean_half_width_sd=round(float(widths.mean()), 3)),
                      totals=totals(rows), includes_highest=totals(top) if top else None)


def totals(rows):
    n = len(rows)
    rep3 = sum(r['next_step'] == 'report at three' for r in rows)
    rep4 = sum(r['next_step'] == 'report after one more' for r in rows)
    used, full = sum(r['wells_used'] for r in rows), sum(r['wells_full'] for r in rows)
    mo_used = sum(r['measured_only_wells'] for r in rows)
    ab_used = sum(r['two_action_wells'] for r in rows)
    return dict(designs=n, compounds=len({r['compound'] for r in rows}), calls=rep3 + rep4, calls_from_three=rep3,
                calls_after_one_more=rep4, completed=n - rep3 - rep4, wrong_calls=sum(r['wrong'] for r in rows),
                wells_used=used, wells_full=full, wells_saved_vs_full=round(1 - used / full, 4),
                measured_only=dict(calls=sum(r['measured_only_report'] for r in rows),
                                   wrong_calls=sum(r['measured_only_wrong'] for r in rows), wells_used=mo_used),
                two_action=dict(calls=sum(r['two_action_report'] for r in rows),
                                wrong_calls=sum(r['two_action_wrong'] for r in rows), wells_used=ab_used),
                wells_saved_vs_measured_only=round(1 - used / mo_used, 4))


def summary_replay(per, args, digest_lines):
    lines = [f'# Replay of three-concentration designs from {Path(args.table).name}', '',
             'Each completed series was replayed from five three-concentration designs. Compounds were dealt into five '
             'folds; for each test fold, two folds trained the model and two calibrated the release margin, so no test '
             'compound informed its own forecast or call. Calls are compared with the call from the full series.', '',
             '| Endpoint | Compounds | Designs | Calls from three | After one more | Wrong calls | Wells used | Saved against the full series | Measured concentrations only: calls, wrong | Saved against measured-only |',
             '|:--|--:|--:|--:|--:|--:|--:|--:|--:|--:|']
    for ep, p in sorted(per.items()):
        t = p['totals']
        lines.append(f'| {ep} | {t["compounds"]} | {t["designs"]} | {t["calls_from_three"]} | {t["calls_after_one_more"]} | '
                     f'{t["wrong_calls"]} | {t["wells_used"]:,} of {t["wells_full"]:,} | {t["wells_saved_vs_full"]:.1%} | '
                     f'{t["measured_only"]["calls"]}, {t["measured_only"]["wrong_calls"]} | {t["wells_saved_vs_measured_only"]:.1%} |')
    lines += ['', '| Endpoint | Forecast error, AnchorBoost | Interpolation | Difference, paired 95% interval | Hidden wells inside the 90% interval |',
              '|:--|--:|--:|--:|--:|']
    for ep, p in sorted(per.items()):
        c, i = p['curve_mae'], p['intervals']
        lines.append(f'| {ep} | {c["anchorboost"]:.3f} | {c["loglinear"]:.3f} | {c["paired"]["mean_diff"]:+.3f} '
                     f'[{c["paired"]["ci95"][0]:+.3f}, {c["paired"]["ci95"][1]:+.3f}] | {i["coverage"]:.1%} of {i["wells"]:,} |')
    lines += ['', f'Forecast error is the mean absolute difference from the measured concentration means, in baseline SDs '
              f'(the spread of responses at the two lowest concentrations of training and calibration compounds, after '
              f'centring each plate on its vehicle wells when the table has them), per compound over its designs. A call is active when the largest '
              f'absolute concentration mean reaches {args.threshold:g} baseline SDs; the release margin keeps the expected share of wrong early '
              f'calls per compound at or below {args.alpha:.0%} on the calibration compounds. "Measured concentrations '
              'only" is the same calibrated rule reading the three measured concentrations without a forecast. Wells '
              'count exposed wells; vehicle wells are shared and not counted.', '', '**Inputs and code (SHA-256)**', '']
    lines += [f'- {k}: `{v}`' for k, v in digest_lines]
    return '\n'.join(lines) + '\n'


# ---------------------------------------------------------------- reconcile

def reconcile(args):
    start = time.time()
    lockdir = Path(args.reconcile)
    lock = json.loads((lockdir / 'lock.json').read_text())
    for name, digest in [('forecast.csv', lock['output']['forecast.csv'])] + \
            [(m['file'], m['sha256']) for m in lock['models'].values()]:
        if sha(lockdir / name) != digest:
            raise ValueError(f'{lockdir / name} differs from the locked SHA-256')
    with open(lockdir / 'forecast.csv', newline='') as f:
        locked = list(csv.DictReader(f))
    cf.ND, cf.NF, cf.D = 1, 1, 1
    finished = {(c.compound, c.endpoint): c for c in load_curves(args.table)}
    bundles = {}
    for ep, m in lock['models'].items():
        with open(lockdir / m['file'], 'rb') as fh:
            bundles[ep] = pickle.load(fh)['fit']
    by = defaultdict(list)
    for r in locked:
        by[(r['compound'], r['endpoint'])].append(r)
    rows, per_curve = [], []
    for key, rr in sorted(by.items()):
        if key not in finished:
            per_curve.append(dict(compound=key[0], endpoint=key[1], status='not in the finished table'))
            continue
        c, f = finished[key], bundles[key[1]]
        t = c.task(f['scale'])
        y = truth(t, f['threshold'])
        hits = wells = 0
        for r in rr:
            if r['status'] != 'forecast':
                continue
            lv = np.float32(round(np.log10(float(r['concentration'])), 4))
            i = np.argmin(np.abs(t.levels - lv))
            if abs(float(t.levels[i]) - float(lv)) > TOL:
                rows.append(dict(compound=key[0], endpoint=key[1], concentration=r['concentration'], unit=r['unit'],
                                 status='not measured in the finished table'))
                continue
            vals = t.y[t.logc == t.levels[i]].reshape(-1)
            lo, hi = float(r['response_low90']), float(r['response_high90'])
            inside = int(((vals >= lo) & (vals <= hi)).sum())
            hits, wells = hits + inside, wells + len(vals)
            mean = float(t.mu[i].ravel()[0])
            rows.append(dict(compound=key[0], endpoint=key[1], concentration=r['concentration'], unit=r['unit'],
                             status='measured', forecast=r['value'], low90=r['low90'], high90=r['high90'],
                             measured_mean=round(c.measured_mean(t.levels[i]), 6), wells=len(vals), wells_inside_90=inside,
                             response_forecast=r['response'], response_measured=round(mean, 4),
                             abs_error_sd=round(abs(mean - float(r['response'])), 4)))
        step, call = rr[0]['next_step'], rr[0]['call']
        measured = sorted({np.float32(round(np.log10(float(r['concentration'])), 4)) for r in rr if r['status'] == 'measured'})
        ctx = np.array([t.levels[np.argmin(np.abs(t.levels - v))] for v in measured], np.float32)
        wells_initial = int(np.isin(t.logc, ctx).sum())
        outcome = dict(compound=key[0], endpoint=key[1], next_step=step, full_series_call='inactive' if y else 'active',
                       wells_inside_90=hits, hidden_wells=wells, wells_full=len(t.logc))
        if step == 'report the call':
            outcome.update(call=call, wrong=int((call == 'inactive') != bool(y)), wells_used=wells_initial)
        elif step == 'add one concentration':
            add = np.float32(round(np.log10(float(rr[0]['add_concentration'])), 4))
            level = t.levels[np.argmin(np.abs(t.levels - add))]
            ctx4 = np.sort(np.r_[ctx, level])
            s4 = score(f['model'], observed(t, ctx4), t.levels[~np.isin(t.levels, ctx4)])[0]
            call4, margin4 = nd.apply_rule(f['rules']['methods']['anchorboost'], s4)
            added = int(np.sum(t.logc == level))
            if margin4 > f['chain_margin']:
                c4 = 'inactive' if call4 else 'active'
                outcome.update(after_one_more='report the call', call=c4, wrong=int(call4 != y),
                               wells_used=wells_initial + added)
            else:
                outcome.update(after_one_more='complete the series', call='', wrong=0, wells_used=len(t.logc))
        else:
            outcome.update(call='', wrong=0, wells_used=len(t.logc))
        per_curve.append(outcome)
    out = Path(args.out) if args.out else lockdir
    write_rows(out / 'reconcile.csv', rows)
    done = [p for p in per_curve if 'wells_used' in p]
    reported = [p for p in done if p.get('call')]
    total = dict(series=len(done), calls=len(reported), wrong_calls=sum(p['wrong'] for p in reported),
                 completed=len(done) - len(reported), hidden_wells=sum(p['hidden_wells'] for p in done),
                 wells_inside_90=sum(p['wells_inside_90'] for p in done), wells_used=sum(p['wells_used'] for p in done),
                 wells_full=sum(p['wells_full'] for p in done))
    lines = [f'# Locked forecast against the finished series', '',
             f'Lock `{sha(lockdir / "lock.json")}` from {lock["created_utc"]} (UTC); finished table '
             f'{Path(args.table).name}, SHA-256 `{sha(args.table)}`. Every locked file matched its SHA-256.', '',
             f'**{total["wells_inside_90"]:,} of {total["hidden_wells"]:,} hidden wells** fell inside the locked 90% intervals. '
             f'**{total["calls"]} of {total["series"]} series** were called before the series was finished, '
             f'**{total["wrong_calls"]} of them wrong** against the full-series call; {total["completed"]} were completed. '
             f'The steps used {total["wells_used"]:,} of {total["wells_full"]:,} exposed wells.', '',
             '| Compound | Endpoint | Next step | After one more | Call | Full series | Hidden wells inside 90% |',
             '|:--|:--|:--|:--|:--|:--|--:|']
    for p in per_curve:
        if 'wells_used' not in p:
            lines.append(f'| {p["compound"]} | {p["endpoint"]} | {p["status"]} | | | | |')
            continue
        lines.append(f'| {p["compound"]} | {p["endpoint"]} | {p["next_step"]} | {p.get("after_one_more", "")} | '
                     f'{p.get("call", "")}{" (wrong)" if p["wrong"] else ""} | {p["full_series_call"]} | '
                     f'{p["wells_inside_90"]} of {p["hidden_wells"]} |')
    (out / 'reconcile.md').write_text('\n'.join(lines) + '\n')
    total['runtime'] = runtime(start)
    (out / 'reconcile.json').write_text(json.dumps(dict(lock_sha256=sha(lockdir / 'lock.json'), finished_table=Path(args.table).name,
                                                        finished_table_sha256=sha(args.table), totals=total,
                                                        series=per_curve), indent=1, default=float) + '\n')
    print(json.dumps(dict(mode='reconcile', output=str(out), **total), indent=1))


# ---------------------------------------------------------------- example table

def example(out):
    """The US EPA developmental-neurotoxicity screen in human neural cells (Harrill et al. 2018, public domain), as a
    long table: the 9 human-cell endpoints, accepted wells, vehicle (DMSO) wells as concentration 0. A second table
    keeps three concentrations (lowest, middle, highest) of one compound in five and leaves the rest to measure."""
    import openpyxl
    import chip_forecast_dnt as dnt
    path = dnt.fetch(ROOT / '.cache' / 'epa_dnt_hci')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = {r[0]: r[1] for r in list(wb['assay_component'].iter_rows(values_only=True))[1:]}
    rows = []
    for r in list(wb['Data'].iter_rows(values_only=True))[1:]:
        ep = names[r[1]]
        if ep.split('_')[0] not in dnt.HUMAN or r[7] != 1 or r[6] not in ('t', 'n') or r[9] is None:
            continue
        vehicle = r[6] == 'n'
        rows.append(dict(compound='DMSO' if vehicle else str(r[2]).strip(), concentration=0 if vehicle else r[8],
                         unit='uM', endpoint=ep, value=r[9], plate=r[3], date=''))
    out.mkdir(parents=True, exist_ok=True)
    full = out / 'epa_dnt_human_neural.csv'
    write_rows(full, rows)
    chems = sorted({r['compound'] for r in rows if r['concentration']})
    new = {c for c, g in deal(chems, ['example-v1'], 5).items() if g == 0}
    levels = defaultdict(set)
    for r in rows:
        if r['compound'] in new and r['concentration']:
            levels[(r['compound'], r['endpoint'])].add(r['concentration'])
    keep = {k: (lambda v: {v[0], v[(len(v) - 1) // 2], v[-1]})(sorted(v)) for k, v in levels.items()}
    three = [r for r in rows if r['compound'] not in new or not r['concentration']
             or r['concentration'] in keep[(r['compound'], r['endpoint'])]]
    for (c, ep), v in sorted(levels.items()):
        three += [dict(compound=c, concentration=x, unit='uM', endpoint=ep, value='', plate='', date='')
                  for x in sorted(v - keep[(c, ep)])]
    part = out / 'epa_dnt_three_point.csv'
    write_rows(part, three)
    print(json.dumps(dict(source=dnt.URL, source_sha256=dnt.SHA256, full_table=str(full), rows=len(rows),
                          compounds=len(chems), three_point_table=str(part), compounds_with_three_points=sorted(new)), indent=1))


# ---------------------------------------------------------------- shared output

def write_rows(path, rows):
    if not rows:
        raise ValueError(f'nothing to write to {path}')
    fields = list(dict.fromkeys(k for r in rows for k in r))
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def protocol(args, mode, files):
    return dict(schema='three_point.protocol.v1', mode=mode, threshold=args.threshold, alpha=args.alpha,
                measured_concentrations=K, margin_grid=nd.MARGIN_GRID, concentration_level_tolerance_log10=TOL,
                response_scale='value minus the median of its plate\'s vehicle wells (when the table has vehicle '
                               'wells for the endpoint), then minus the baseline median and divided by the baseline '
                               '1.4826 x MAD, both from the two lowest concentrations of every training and calibration '
                               'series (tcpl), clipped to +-10',
                call='inactive when the largest absolute concentration mean of the full series is below the threshold',
                model=dict(name='AnchorBoost', params=cf.MODEL, designs='all C(L,3) training designs'),
                rule='paper_nested_decision: training-only cutoff and weighted rank margin from two inner fits, exact '
                     'margins; joint chain margin over every candidate fourth concentration; group = compound',
                intervals=dict(method='chip_forecast_intervals learned width, cross-fitted over calibration parts',
                               parts=PARTS, params=ci.WIDTH, rows=ci.ROWS, floor=ci.FLOOR, level=sc.LEVEL),
                partition=('forecast: completed compounds split half and half into training and calibration by '
                           'SHA-256; replay: five folds, two training, two calibration, one test'),
                inputs=files, code={n: sha(ROOT / n) for n in CODE},
                versions=dict(python=platform.python_version(), numpy=np.__version__, sklearn=sklearn.__version__,
                              platform=platform.platform()))


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('table', nargs='?', type=Path, help='long CSV: compound, concentration, unit, endpoint, value, plate, date')
    p.add_argument('--library', type=Path, nargs='+', default=[], help='more tables of completed series')
    p.add_argument('--endpoint', nargs='+', help='use only these endpoints')
    p.add_argument('--mode', choices=['auto', 'forecast', 'replay'], default='auto')
    p.add_argument('--threshold', type=float, default=3.0, help='activity threshold in baseline SDs (default 3)')
    p.add_argument('--alpha', type=float, default=nd.ALPHA, help='expected wrong early calls per compound (default 0.10)')
    p.add_argument('--out', type=Path, help='output folder (default three_point_<table name>)')
    p.add_argument('--jobs', type=int, default=os.cpu_count() or 1,
                   help='worker processes, one endpoint or fold each (default: one per CPU core)')
    p.add_argument('--reconcile', type=Path, metavar='LOCKED_FOLDER', help='check this finished table against a locked forecast')
    p.add_argument('--example', type=Path, metavar='FOLDER', help='write the EPA human neural screen as example tables')
    args = p.parse_args()
    if args.example:
        example(args.example)
        return
    if args.table is None:
        p.error('give a table')
    if args.reconcile:
        reconcile(args)
        return
    cf.ND, cf.NF, cf.D = 1, 1, 1  # one output per endpoint series
    start = time.time()
    curves = load_curves(args.table, args.endpoint)
    library = [c for path in args.library for c in load_curves(path, args.endpoint)]
    open_curves = [c for c in curves if c.open]
    mode = args.mode if args.mode != 'auto' else 'forecast' if open_curves else 'replay'
    out = args.out or Path(f'three_point_{args.table.stem}')
    out.mkdir(parents=True, exist_ok=True)
    files = {Path(x).name: sha(x) for x in [args.table, *args.library]}
    completed = [c for c in curves if not c.open] + library
    seen = {}
    for c in completed:
        if (c.compound, c.endpoint) in seen:
            raise ValueError(f'{c.compound} / {c.endpoint} is a completed series in both {seen[(c.compound, c.endpoint)]} and {c.source}')
        seen[(c.compound, c.endpoint)] = c.source
    if mode == 'forecast':
        if not open_curves:
            raise ValueError('no series with concentrations still to measure (rows with an empty value)')
        (out / 'protocol.json').write_text(json.dumps(protocol(args, mode, files), indent=1) + '\n')
        eps = sorted({c.endpoint for c in open_curves})
        parts = run_all(args.jobs, forecast_endpoint,
                        [(ep, [c for c in open_curves if c.endpoint == ep], [c for c in completed if c.endpoint == ep],
                          args.threshold, args.alpha, out) for ep in eps])
        table = [r for p_ in parts for r in p_['table']]
        write_rows(out / 'forecast.csv', table)
        models = {p_['endpoint']: p_['model'] for p_ in parts}
        settings = {p_['endpoint']: p_['settings'] for p_ in parts}
        notes = [n for p_ in parts for n in p_['notes']]
        created = datetime.now(timezone.utc).isoformat(timespec='seconds')
        digest_lines = [('input ' + Path(k).name, v) for k, v in files.items()] + \
                       [('forecast.csv', sha(out / 'forecast.csv'))] + \
                       [(f'model {ep} ({m["file"]})', m['sha256']) for ep, m in models.items()] + \
                       [('protocol.json', sha(out / 'protocol.json'))] + \
                       [(f'code {n}', sha(ROOT / n)) for n in CODE]
        (out / 'summary.md').write_text(summary_forecast(table, settings, notes, args,
                                                         dict(created_utc=created, digest_lines=digest_lines)))
        lock = dict(schema='three_point.lock.v1', created_utc=created, inputs=files,
                    output={'forecast.csv': sha(out / 'forecast.csv'), 'summary.md': sha(out / 'summary.md')},
                    models=models, protocol_sha256=sha(out / 'protocol.json'), code={n: sha(ROOT / n) for n in CODE},
                    settings=settings, runtime=runtime(start))
        (out / 'lock.json').write_text(json.dumps(lock, indent=1) + '\n')
        steps = defaultdict(int)
        for step in {(r['compound'], r['endpoint']): r['next_step'] for r in table}.values():
            steps[step] += 1
        print(json.dumps(dict(mode='forecast', series=sum(steps.values()), next_steps=dict(steps), output=str(out),
                              lock_sha256=sha(out / 'lock.json'), runtime=lock['runtime']), indent=1))
        return
    by_ep = defaultdict(list)
    for c in completed:
        if len(c.levels) > K:
            by_ep[c.endpoint].append(c)
    if not by_ep:
        raise ValueError('no completed series with at least four concentrations to replay')
    (out / 'protocol.json').write_text(json.dumps(protocol(args, mode, files), indent=1) + '\n')
    for ep, cc in by_ep.items():
        if len({c.unit for c in cc}) > 1:
            raise ValueError(f'endpoint {ep!r}: concentrations in different units {sorted({c.unit for c in cc})}')
    tasks = [(cc, ep, f, args.threshold, args.alpha) for ep, cc in sorted(by_ep.items()) for f in range(5)
             if any(g == f for g in deal([c.compound for c in cc], ['outer-v1', ep], 5).values())]
    parts = run_all(args.jobs, replay_fold, tasks)
    per, rows = {}, []
    for ep in sorted(by_ep):
        r, info = replay_summary([p_ for p_ in parts if p_['endpoint'] == ep])
        per[ep], rows = info, rows + r
    write_rows(out / 'replay.csv', rows)
    digest_lines = [('input ' + Path(k).name, v) for k, v in files.items()] + [('replay.csv', sha(out / 'replay.csv'))] + \
                   [('protocol.json', sha(out / 'protocol.json'))] + [(f'code {n}', sha(ROOT / n)) for n in CODE]
    (out / 'summary.md').write_text(summary_replay(per, args, digest_lines))
    result = dict(schema='three_point.replay.v1', created_utc=datetime.now(timezone.utc).isoformat(timespec='seconds'),
                  inputs=files, protocol_sha256=sha(out / 'protocol.json'), replay_csv_sha256=sha(out / 'replay.csv'),
                  endpoints=per, all_endpoints=totals(rows), runtime=runtime(start))
    (out / 'summary.json').write_text(json.dumps(result, indent=1, default=float) + '\n')
    print(json.dumps(dict(mode='replay', output=str(out), all_endpoints=result['all_endpoints'],
                          runtime=result['runtime']), indent=1))


if __name__ == '__main__':
    main()
