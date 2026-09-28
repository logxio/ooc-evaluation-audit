#!/usr/bin/env python3
"""Recompute an exploratory assay-discordance and nested channel-selection audit.

The Shtenberg Figure 5 patient-derived spheroid measurements are an adjacent
assay, not confirmed measurements in a perfused chip. No clinical accuracy is
inferred from those rows.
"""

import argparse
import hashlib
import io
import json
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from chip_clinic import (
    FIGURE_SHA256, FIGURE_URL, PATIENTS, SOURCE_SHA256, SOURCE_URL,
    balanced_accuracy, evaluate, read_chip, read_clinical, source_bytes,
    threshold,
)


SPHEROID_URL = ('https://static-content.springer.com/esm/'
                'art%3A10.1038%2Fs42003-023-05531-5/MediaObjects/'
                '42003_2023_5531_MOESM3_ESM.xlsx')
SPHEROID_SHA256 = 'cb4375562f4d3a48a8c5246c921b1dd7277870bf87f404a95386d3376e94f660'
CHANNEL_PRIORITY = (
    'optimized_mean', 'optimized_vessel', 'optimized_tumoroid',
    'original_mean', 'original_vessel', 'original_tumoroid',
)
DRUGS = (
    'Oxaliplatin', '5-FU', 'Oxaliplatin+5-FU', 'Gemcitabine',
    'Cisplatin', 'Etoposide', 'Mitomycin', 'Bevacizumab',
    'Bevacizumab+5-FU',
)


def nested_patient_audit(chip, clinical):
    labels = [clinical[p] for p in PATIENTS]
    scores = {}
    for condition in ('original', 'optimized'):
        for key in ('vessel', 'tumoroid'):
            scores[f'{condition}_{key}'] = [chip[p][key][condition] for p in PATIENTS]
        scores[f'{condition}_mean'] = [
            (chip[p]['vessel'][condition] + chip[p]['tumoroid'][condition]) / 2
            for p in PATIENTS
        ]
    selected, predicted, cuts, inner_ba = [], [], [], []
    for held in range(len(PATIENTS)):
        train = [i for i in range(len(PATIENTS)) if i != held]
        train_truth = [labels[i] for i in train]
        inner = {name: evaluate(train_truth, [scores[name][i] for i in train])['balanced_accuracy']
                 for name in CHANNEL_PRIORITY}
        name = max(CHANNEL_PRIORITY, key=lambda key: inner[key])
        cut = threshold([scores[name][i] for i in train], train_truth)
        selected.append(name)
        predicted.append(int(scores[name][held] <= cut))
        cuts.append(cut)
        inner_ba.append(inner)
    ba, sensitivity, specificity = balanced_accuracy(labels, predicted)
    return {
        'patients': list(PATIENTS), 'clinical_sensitive': labels,
        'selected_channel': selected, 'outer_predictions': predicted,
        'outer_cutoffs': cuts, 'inner_balanced_accuracy': inner_ba,
        'channel_counts': dict(sorted(Counter(selected).items())),
        'correct': sum(t == p for t, p in zip(labels, predicted)),
        'n': len(labels), 'balanced_accuracy': ba,
        'clinical_sensitive_recall': sensitivity,
        'clinical_resistant_recall': specificity,
        'a9_fixed_mean_correct': round(evaluate(labels, scores['optimized_mean'])['accuracy'] * len(labels)),
        'a9_exploratory_vessel_correct': round(evaluate(labels, scores['optimized_vessel'])['accuracy'] * len(labels)),
    }


def discordance_audit(data):
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook['Fig. 5']
    assert tuple(sheet.cell(2, col).value for col in range(3, 13)) == tuple(range(1, 11))
    assert tuple(sheet.cell(13, col).value for col in range(3, 13)) == tuple(range(1, 11))
    pairs = []
    for index, drug in enumerate(DRUGS):
        assert sheet.cell(3 + index, 2).value == drug
        assert sheet.cell(14 + index, 2).value == drug
        for patient in range(1, 11):
            viability = sheet.cell(3 + index, patient + 2).value
            area = sheet.cell(14 + index, patient + 2).value
            assert (viability is None) == (area is None)
            if viability is None:
                continue
            assert isinstance(viability, (int, float)) and isinstance(area, (int, float))
            if viability * area < 0:
                category = 'area_expands_viability_falls' if viability > 0 else 'area_shrinks_viability_rises'
            elif viability == 0 or area == 0:
                category = 'neutral'
            elif viability > 0:
                category = 'both_reduced'
            else:
                category = 'both_increased'
            pairs.append({'patient': patient, 'drug': drug, 'viability_reduction': viability,
                          'area_reduction': area, 'category': category})
    by_drug = defaultdict(list)
    for pair in pairs:
        by_drug[pair['drug']].append(pair)
    counts = Counter(p['category'] for p in pairs)
    assert len(pairs) == 49
    viability_reduced = counts['both_reduced'] + counts['area_expands_viability_falls']
    return {
        'n_pairs': len(pairs), 'n_patients': 10,
        'article_figure_caption_n': 48,
        'caption_source_table_denominator_discrepancy': 'Figure 5 caption says n=48; 49 cells have both numeric readouts in Supplementary Data 1 Figure 5 rows 3-11 and 14-22.',
        'category_counts': dict(sorted(counts.items())),
        'discordant': counts['area_expands_viability_falls'] + counts['area_shrinks_viability_rises'],
        'discordance_rate': (counts['area_expands_viability_falls'] + counts['area_shrinks_viability_rises']) / len(pairs),
        'area_only_missed_viability_reductions': counts['area_expands_viability_falls'],
        'viability_reduced_pairs': viability_reduced,
        'area_sign_sensitivity_to_viability_reduction': counts['both_reduced'] / viability_reduced,
        'by_drug': {drug: {'n': len(by_drug[drug]),
                           'discordant': sum(p['category'] in ('area_expands_viability_falls', 'area_shrinks_viability_rises') for p in by_drug[drug]),
                           'area_only_missed_viability_reductions': sum(p['category'] == 'area_expands_viability_falls' for p in by_drug[drug])}
                    for drug in DRUGS},
        'pairs': pairs,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-zip', type=Path)
    parser.add_argument('--figure', type=Path)
    parser.add_argument('--spheroid', type=Path)
    parser.add_argument('--out', type=Path, help='Optional private patient-level audit JSON')
    args = parser.parse_args()
    chip = read_chip(source_bytes(args.source_zip, SOURCE_URL, SOURCE_SHA256))
    clinical = read_clinical(source_bytes(args.figure, FIGURE_URL, FIGURE_SHA256))
    spheroid = source_bytes(args.spheroid, SPHEROID_URL, SPHEROID_SHA256)
    patient_result = nested_patient_audit(chip, clinical)
    spheroid_result = discordance_audit(spheroid)
    result = {
        'schema': 'assay.discordance.v1',
        'dai_source_sha256': SOURCE_SHA256, 'dai_figure_sha256': FIGURE_SHA256,
        'shtenberg_source_url': SPHEROID_URL, 'shtenberg_source_sha256': SPHEROID_SHA256,
        'shtenberg_license': 'CC BY 4.0',
        'channel_priority': list(CHANNEL_PRIORITY),
        'nested_patient': patient_result, 'spheroid_assay': spheroid_result,
        'limits': ['Nested algorithm conceived after inspecting six A9 readouts; exploratory, no independent patient cohort.',
                   'Shtenberg Fig.5 does not explicitly establish perfused-chip measurements for patient spheroids; adjacent assay only.',
                   'Treatment and sampling time do not support pooled patient-level clinical accuracy from the spheroid table.'],
    }
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(result, indent=2) + '\n')
    compact = {
        'nested_patient': {key: patient_result[key] for key in ('n', 'correct', 'balanced_accuracy',
            'clinical_sensitive_recall', 'clinical_resistant_recall', 'channel_counts',
            'a9_fixed_mean_correct', 'a9_exploratory_vessel_correct')},
        'spheroid_assay': {key: spheroid_result[key] for key in ('n_pairs', 'discordant',
            'discordance_rate', 'area_only_missed_viability_reductions', 'viability_reduced_pairs',
            'area_sign_sensitivity_to_viability_reduction', 'category_counts', 'by_drug')},
    }
    print(json.dumps(compact, indent=2))


if __name__ == '__main__':
    main()
