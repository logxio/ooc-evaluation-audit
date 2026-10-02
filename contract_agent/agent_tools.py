"""Source browsing, visual evidence and short native-sandbox calculations."""
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import threading
import time

from .engine import evaluate

HERE=Path(__file__).resolve().parent


def load(path):return json.loads(Path(path).read_text())


def function(name,description,properties,required=()):
    return {'type':'function','function':{'name':name,'description':description,'parameters':
            {'type':'object','properties':properties,'required':list(required),'additionalProperties':False}}}


CROP={'type':'array','items':{'type':'number'},'description':'Optional [left, top, right, bottom] fractions of the rendered image.'}
TOOLS=[
 function('list_files','List every available source file and table sheet, including acquisition errors.',{}),
 function('read_article','Search full article paragraphs, or return a range; exact source references included.',
          {'query':{'type':'string'},'start':{'type':'integer'},'limit':{'type':'integer'}}),
 function('preview_table','Read structured rows and cell references; PDF words include x coordinates; styles preserve workbook fills.',
          {'file':{'type':'string'},'sheet':{'type':'string'},'start':{'type':'integer'},'rows':{'type':'integer'}},['file','sheet']),
 function('view_source','See an original figure, PDF page, or colored workbook rows; page numbers are one-based.',
          {'file':{'type':'string'},'page':{'type':'integer'},'sheet':{'type':'string'},'start':{'type':'integer'},'rows':{'type':'integer'},
           'start_column':{'type':'integer'},'columns':{'type':'integer'},'crop':CROP},['file']),
 function('read_figure','An independent reader that never sees this task, its target or your notes transcribes a figure, PDF page or crop: printed row labels, legend categories of coloured markers or text, approximate axis values. Returns vision/<image>/R<row>C<col> references usable in a contract; you cannot write vision values yourself.',
          {'file':{'type':'string'},'page':{'type':'integer'},'crop':CROP},['file']),
 function('read_grid','A deterministic pixel reader, no model involved, for coloured grids, heat maps and bar charts. Give the grid box and its rows and columns; it samples the centre of every cell and names its colour (red, pink, orange, yellow, olive, green, blue, light blue, lavender, white, grey, black, other) with the RGB it measured. Returns grid/<reading>/R<row>C<col> references usable in a contract and an image of the grid with every sample point marked, so you can check the alignment and adjust the box. snap=true (default) finds the filled cell block inside the box, so a box slightly larger than a heat map is fine; snap=false uses the box exactly, e.g. a thin strip across all bars at one axis value, where a cell reads the bar colour if the bar reaches it and otherwise whatever lies there. You cannot change its labels; translate colours to categories with a map that follows the figure legend.',
          {'file':{'type':'string'},'page':{'type':'integer'},'box':{**CROP,'description':'[left, top, right, bottom] fractions of the whole figure or PDF page, not of a crop'},
           'rows':{'type':'integer'},'cols':{'type':'integer'},'snap':{'type':'boolean'}},['file','box','rows','cols']),
 function('python','Run short Python with full cells, tables, styles, article and files already loaded. Set CONTRACT to a complete executable contract to check it. Output is cached. 20s limit, source read-only, no network, writes confined to scratch.',
          {'code':{'type':'string'}},['code']),
 function('finish','Return a contract or an explicit no-headline/blocked conclusion.',
          {'contract':{'type':'object'}},['contract'])
]


class Workbench:
    def __init__(self,source,scratch,paper,reader=None):
        self.source=Path(source).resolve();self.scratch=Path(scratch).resolve();self.paper=paper;self.reader=reader
        self.scratch.mkdir(parents=True,exist_ok=True)
        self.tables=load(self.source/'tables.json');self.cells=load(self.source/'cells.json')
        self.styles=load(self.source/'styles.json');self.article=load(self.source/'article.json')
        self.files=load(self.source/'inventory.json')
        self.visual_path=self.scratch/'visual.json'
        self.visual=load(self.visual_path) if self.visual_path.exists() else {}
        self.executions=max([int(p.name[6:]) for p in self.scratch.glob('python[0-9]*') if p.name[6:].isdigit()],default=0)

    def index(self):return {**self.cells,**{k:v['value'] for k,v in self.visual.items()}}

    def redact(self,text):return text.replace(str(self.source),'source').replace(str(self.scratch),'scratch')

    def check(self,contract,target):
        if not isinstance(contract,dict):return {'matched':False,'errors':['Contract must be an object']}
        if target is None:
            no_headline=contract.get('status')=='no_headline' and not contract.get('headline',{}).get('records')
            return {'matched':False,'correct_refusal':no_headline,'errors':[] if no_headline else ['Expected an explicit no_headline conclusion']}
        t={'metric':target['metric'],'reported':target['value'],'n':target['n'],'tolerance':target['tolerance']}
        try:
            # Published rounded values have a tolerance; the frozen judge supplies
            # that same tolerance, independently of any model-chosen tolerance.
            h=contract.get('headline',{})
            t['reported']=h.get('reported') if isinstance(h.get('reported'),(int,float)) and abs(h['reported']-target['value'])<=t['tolerance'] else target['value']
            result=evaluate(contract,self.index(),t)
        except Exception as exc:return {'matched':False,'errors':[f'{type(exc).__name__}: {exc}']}
        errors=result['errors']
        if self.paper=='ewart' and contract.get('headline',{}).get('threshold')==375:
            raw=json.dumps(contract.get('headline',{}).get('records',[]))
            if 'article.xml/Table 4/' in raw:errors.append('Table 4 MOS is uncorrected; 375 requires protein-binding correction, not uncorrected Table 4 values.')
        used=json.dumps(contract.get('headline',{}).get('records',[]))
        for ref,item in self.visual.items():
            if json.dumps(ref) in used and item['uncertain']:errors.append(f'Uncertain visual label {ref}')
        result['matched']=result['matched'] and not errors
        result['expected']={k:target[k] for k in ['metric','value','n']}
        return result

    def preview(self,file,sheet,start=1,rows=12):
        data=self.tables[file][sheet];rows=max(1,min(rows,100));start=max(1,start)
        output=[]
        for r,values in enumerate(data.get('rows',[])[start-1:start-1+rows],start):
            row=[]
            for c,value in enumerate(values,1):
                if file.endswith('.xlsx'):
                    from openpyxl.utils import get_column_letter
                    suffix=f'{get_column_letter(c)}{r}'
                else:suffix=f'R{r}C{c}'
                ref=f'{file}/{sheet}/{suffix}'
                if value is not None and value!='':
                    cell={'ref':ref,'value':value}
                    if ref in self.styles:cell['style']=self.styles[ref]
                    if data.get('geometry'):cell['x0']=data['geometry'][r-1]['cells'][c-1]['x0']
                    row.append(cell)
            output.append({'row':r,'cells':row})
        return {'file':file,'sheet':sheet,**({'caption':data['caption']} if data.get('caption') else {}),'nrows':data.get('nrows'),'ncols':data.get('ncols'),
                'rows':output,'merged':data.get('merged',[]),**({'error':data['error']} if 'error' in data else {})}

    def picture(self,args):
        from PIL import Image,ImageDraw,ImageFont
        file=args['file'];path=(self.source/file).resolve()
        if path.parent!=self.source or not path.is_file():raise ValueError('Unknown source file')
        if path.suffix=='.pdf':
            import pypdfium2
            pdf=pypdfium2.PdfDocument(path);page=max(1,int(args.get('page',1)))
            image=pdf[page-1].render(scale=2).to_pil().convert('RGB')
            descriptor=f'{file}, page {page}'
        elif path.suffix in ['.xlsx','.xls']:
            sheet=args.get('sheet') or next(iter(self.tables[file]));start=max(1,args.get('start',1));count=min(30,args.get('rows',16))
            rows=self.tables[file][sheet]['rows'][start-1:start-1+count]
            first=max(1,args.get('start_column',1));cols=max(1,min(14,args.get('columns',14),max(map(len,rows),default=1)-first+1));width=220
            image=Image.new('RGB',(80+cols*width,45+len(rows)*70),'white');draw=ImageDraw.Draw(image)
            try:font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',15)
            except OSError:font=ImageFont.load_default()
            draw.text((10,8),f'{file} / {sheet}',fill='black',font=font)
            from openpyxl.utils import get_column_letter
            for i,row in enumerate(rows):
                y=40+i*70;draw.text((3,y+5),str(start+i),fill='black',font=font)
                for c,value in enumerate(row[first-1:first-1+cols],first):
                    ref=f'{file}/{sheet}/{get_column_letter(c)}{start+i}' if file.endswith('.xlsx') else f'{file}/{sheet}/R{start+i}C{c}'
                    fill=self.styles.get(ref,{});color='white'
                    if fill.get('resolved_rgb'):color='#'+fill['resolved_rgb']
                    elif fill.get('color_type')=='rgb' and len(fill['color'])>=6:color='#'+fill['color'][-6:]
                    elif fill.get('color_type')=='indexed':
                        from openpyxl.styles.colors import COLOR_INDEXED
                        n=int(fill['color']);color='#'+COLOR_INDEXED[n][-6:] if n<len(COLOR_INDEXED) else '#dddddd'
                    elif fill:color='#dddddd'
                    x=80+(c-first)*width;draw.rectangle((x,y,x+width,y+70),fill=color,outline='#aaaaaa')
                    s='' if value is None else str(value);wrapped='\n'.join(s[j:j+26] for j in range(0,min(len(s),104),26))
                    draw.text((x+4,y+3),wrapped,fill='black',font=font)
            descriptor=f'{file}, sheet {sheet}, rows {start}-{start+len(rows)-1}, columns {first}-{first+cols-1}; resolved theme/indexed/RGB fills; text may wrap'
        else:image=Image.open(path).convert('RGB');descriptor=file
        if args.get('crop'):
            left,top,right,bottom=[min(1.0,max(0.0,float(x))) for x in args['crop']]
            if right<=left or bottom<=top:raise ValueError('Crop needs left<right and top<bottom as fractions')
            image=image.crop((int(left*image.width),int(top*image.height),int(right*image.width),int(bottom*image.height)))
            descriptor+=f', crop {[round(x,3) for x in (left,top,right,bottom)]}'
        image.thumbnail((2400,2400))
        data=io.BytesIO();image.save(data,format='PNG');raw=data.getvalue();digest=hashlib.sha256(raw).hexdigest()
        (self.scratch/(digest+'.png')).write_bytes(raw)
        return {'source':descriptor,'image_sha256':digest,'width':image.width,'height':image.height},raw

    def read_figure(self,args):
        if self.reader is None:raise RuntimeError('No independent figure reader is configured for this run')
        meta,raw=self.picture(args);reading=self.reader(raw)
        uncertain={(x[0],x[1]) for x in reading.get('uncertain') or [] if isinstance(x,list) and len(x)==2}
        stem='vision/'+meta['image_sha256'][:12];rows=[]
        for r,row in enumerate(reading.get('rows') or [],1):
            cells=[]
            for c,value in enumerate(row if isinstance(row,list) else [row],1):
                if value in (None,''):continue
                ref=f'{stem}/R{r}C{c}';flag=(r,c) in uncertain
                self.visual[ref]={'value':value,'uncertain':flag,'image_sha256':meta['image_sha256'],'source':meta['source'],'reader':'independent'}
                cells.append({'ref':ref,'value':value,**({'uncertain':True} if flag else {})})
            rows.append({'row':r,'cells':cells})
        self.visual_path.write_text(json.dumps(self.visual,indent=2))
        return {'source':meta['source'],'image_sha256':meta['image_sha256'],'columns':reading.get('columns'),'legend':reading.get('legend'),
                'rows':rows,'notes':reading.get('notes')}

    def read_grid(self,args):
        from .grid_colors import grid,load,overlay
        file=args['file'];path=(self.source/file).resolve()
        if path.parent!=self.source or not path.is_file():raise ValueError('Unknown source file')
        if path.suffix.lower() not in ('.pdf','.png','.jpg','.jpeg','.tif','.tiff','.gif','.bmp','.webp'):raise ValueError('read_grid reads figure images and PDF pages')
        page=max(1,int(args.get('page') or 1)) if path.suffix.lower()=='.pdf' else None
        box=[float(x) for x in args['box']];rows=int(args['rows']);cols=int(args['cols']);snap=bool(args.get('snap',True))
        if len(box)!=4:raise ValueError('box needs four fractions [left, top, right, bottom]')
        image=load(path,page);result=grid(image,box,rows,cols,snap)
        data=io.BytesIO();overlay(image,result).save(data,format='PNG');raw=data.getvalue();digest=hashlib.sha256(raw).hexdigest()
        (self.scratch/(digest+'.png')).write_bytes(raw)
        source=f'{file}{f", page {page}" if page else ""}, box {[round(x,4) for x in box]}, {rows}x{cols}, {"snapped" if snap else "exact"}'
        reading=hashlib.sha256(json.dumps({'file':hashlib.sha256(path.read_bytes()).hexdigest(),'page':page,'box':box,'rows':rows,'cols':cols,'snap':snap}).encode()).hexdigest()
        stem='grid/'+reading[:12];labels=[];measured=[];uncertain=[]
        for r,row in enumerate(result['cells'],1):
            for c,cell in enumerate(row,1):
                self.visual[f'{stem}/R{r}C{c}']={'value':cell['colour'],'uncertain':cell['colour']=='other','image_sha256':digest,'source':source,
                                                 'reader':'pixel','rgb':list(cell['rgb']),'xy':cell['xy']}
                if cell['colour']=='other':uncertain.append([r,c])
            labels.append([cell['colour'] for cell in row]);measured.append(['%02x%02x%02x'%tuple(cell['rgb']) for cell in row])
        self.visual_path.write_text(json.dumps(self.visual,indent=2))
        W,H=image.size;left,top,right,bottom=result['grid_px']
        return {'source':source,'references':f'{stem}/R<row>C<col>, one-based','grid_box':[round(left/W,4),round(top/H,4),round((right+1)/W,4),round((bottom+1)/H,4)],
                'cell_px':result['cell_px'],'image_sha256':digest,'colours':labels,'rgb_hex':measured,'uncertain':uncertain},raw

    def python(self,code):
        self.executions+=1
        folder=self.scratch/f'python{self.executions:02d}';folder.mkdir(exist_ok=True)
        (self.scratch/'visual_values.json').write_text(json.dumps({k:v['value'] for k,v in self.visual.items()}))
        wrapper='''import json,math,statistics,re,sys,traceback,resource
from pathlib import Path
import numpy as np
resource.setrlimit(resource.RLIMIT_CPU,(20,20))
resource.setrlimit(resource.RLIMIT_FSIZE,(20000000,20000000))
root=Path(sys.argv[1]); scratch=Path(sys.argv[2])
cells=json.loads((root/'cells.json').read_text());tables=json.loads((root/'tables.json').read_text())
styles=json.loads((root/'styles.json').read_text());article=json.loads((root/'article.json').read_text());files=json.loads((root/'inventory.json').read_text())
cells.update(json.loads((scratch.parent/'visual_values.json').read_text()))
def cell_ref(file,sheet,row,col):
    if file.endswith('.xlsx'):
        letters=''
        while col:col,r=divmod(col-1,26);letters=chr(65+r)+letters
        return f'{file}/{sheet}/{letters}{row}'
    return f'{file}/{sheet}/R{row}C{col}'
def get_rows(file,sheet):return tables[file][sheet]['rows']
def get_cells(prefix):return {k:v for k,v in cells.items() if k.startswith(prefix)}
code=(scratch/'code.py').read_text()
exec(compile(code,'<agent-code>','exec'))
if 'CONTRACT' in globals():
    (scratch/'contract.json').write_text(json.dumps(CONTRACT,ensure_ascii=False,allow_nan=False))
'''
        (folder/'code.py').write_text(code);(folder/'worker.py').write_text(wrapper)
        if sys.platform!='darwin':raise RuntimeError('This live-code adapter needs macOS sandbox-exec; cached replay is cross-platform.')
        # macOS system/runtime files are readable; all user-home trees are denied
        # except this source and scratch. Network and writes outside scratch denied.
        profile='(version 1)\n(allow default)\n(deny network*)\n(deny file-read* (subpath "/Users"))\n'
        for path in [self.source,self.scratch]:profile+='(allow file-read* (subpath '+json.dumps(str(path))+'))\n'
        profile+='(deny file-write*)\n(allow file-write* (subpath '+json.dumps(str(folder))+'))\n(deny process-fork)\n'
        (folder/'profile.sb').write_text(profile)
        env={k:v for k,v in os.environ.items() if k in ['PATH','LANG','LC_ALL','PYTHONDONTWRITEBYTECODE'] or k.startswith(('CODEX_','COMP_JOB_','TASK_WRITER_'))}
        env.update(OPENBLAS_NUM_THREADS='1',OMP_NUM_THREADS='1',PYTHONDONTWRITEBYTECODE='1')
        command=['/usr/bin/sandbox-exec','-f',str(folder/'profile.sb'),sys.executable,'-I','-B',str(folder/'worker.py'),str(self.source),str(folder)]
        started=time.monotonic();peak=0;reason=None
        with (folder/'stdout.txt').open('w') as stdout,(folder/'stderr.txt').open('w') as stderr:
            process=subprocess.Popen(command,cwd=folder,env=env,stdout=stdout,stderr=stderr,start_new_session=True)
            while process.poll() is None:
                try:
                    rss=int(subprocess.check_output(['ps','-o','rss=','-p',str(process.pid)],stderr=subprocess.DEVNULL).strip())*1024
                    peak=max(peak,rss)
                    if rss>1_500_000_000:reason='RSS limit exceeded';process.kill()
                except (ValueError,subprocess.CalledProcessError):pass
                if time.monotonic()-started>20:reason='20 second execution limit exceeded';process.kill()
                time.sleep(.05)
            process.wait()
        def clean(path):return self.redact(path.read_text(errors='replace')[:24000])
        result={'returncode':process.returncode,'stdout':clean(folder/'stdout.txt'),'stderr':clean(folder/'stderr.txt'),
                'seconds':time.monotonic()-started,'peak_rss_bytes':peak,'limit_error':reason,'code_sha256':hashlib.sha256(code.encode()).hexdigest()}
        if (folder/'contract.json').exists():result['contract']=load(folder/'contract.json')
        return result

    def call(self,name,args):
        if name=='list_files':return {'files':self.files},None
        if name=='read_article':
            terms=args.get('query','').lower().split();rows=self.article
            if terms:rows=[p for p in rows if all(t in p['text'].lower() for t in terms)]
            return {'paragraphs':rows[args.get('start',0):args.get('start',0)+min(args.get('limit',8),20)]},None
        if name=='preview_table':return self.preview(**args),None
        if name=='view_source':return self.picture(args)
        if name=='read_figure':return self.read_figure(args),None
        if name=='read_grid':return self.read_grid(args)
        if name=='python':return self.python(args['code']),None
        if name=='finish':return {'contract':args['contract']},None
        raise ValueError('Unknown tool '+name)
